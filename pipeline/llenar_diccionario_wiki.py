"""Llena la hoja «2. Diccionario» del archivo del grupo usando SOLO la wiki y las fuentes F01–F18 del grupo.

Solo entran campos de fuentes con estado «Integrada al lago» en la hoja 1 (F01–F06, F08–F15).
Uso: python pipeline/llenar_diccionario_wiki.py <entrada.xlsx> <salida.xlsx>
"""
from __future__ import annotations

import re
import sys
from copy import copy

import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange

P = "lago/poblacion.json"
V = "lago/vivienda.json"
FI = "lago/finanzas_publicas.json"
S = "lago/seguridad.json"
T = "lago/turismo.json"
SU = "lago/suelo.json"
PQ = "lago/suelo_parques.json"

# (ID, archivo, campo, nombre original, significado, tipo, unidad, ejemplo, vacío, rango, transformación, pregunta)
CAMPOS = [
    # ---------------- Índice 1 — Población ----------------
    ("F01", P, "poblacion_mitad_2026", "Mid-year population (provisional)", "Población estimada de Hong Kong a mediados de 2026.", "Entero", "personas", 7518300, "No", "≥ 0",
     "Se copió la cifra provisional y se quitaron los separadores de miles (7.518.300 → 7518300).", "¿Cuánta gente vive hoy en Hong Kong? (Wiki 1)"),
    ("F01", P, "variacion_poblacion_anual_pct", "Growth rate of population", "Variación de la población frente a mediados de 2025.", "Decimal", "%", 0.3, "No", "−5 a 5",
     "Texto «+0,3 %» convertido a número decimal con punto.", "¿La población está creciendo o bajando? (Wiki 1)"),
    ("F01", P, "densidad_hab_km2", "Population density", "Habitantes por km² de tierra en Hong Kong.", "Entero", "hab/km²", 6801, "No", "≥ 0",
     "Se quitó el «~» de aproximación y los separadores de miles.", "¿Qué tan densamente poblada está la ciudad? (Wiki 1.2)"),
    ("F01", P, "nacimientos_jul2025_jun2026", "Births", "Nacimientos registrados entre mediados de 2025 y mediados de 2026.", "Entero", "nacimientos", 29700, "No", "≥ 0",
     "Se quitó el «~» y se pasó a entero.", "¿Cuántos niños nacen al año? (Wiki 1.10)"),
    ("F01", P, "defunciones_jul2025_jun2026", "Deaths", "Defunciones registradas entre mediados de 2025 y mediados de 2026.", "Entero", "defunciones", 50100, "No", "≥ 0",
     "Se quitó el «~» y se pasó a entero.", "¿Cuántas personas mueren al año? (Wiki 1.10)"),
    ("F01", P, "crecimiento_natural_jul2025_jun2026", "Natural decrease", "Nacimientos menos defunciones en el mismo periodo (negativo = decrecimiento).", "Entero", "personas", -20400, "No", "cualquier entero",
     "Se guardó con signo negativo porque la fuente lo presenta como «déficit natural».", "¿La población crece por nacimientos o solo por migración? (Wiki 1, 1.10)"),
    ("F02", P, "poblacion_proyectada_2046_base", "Projected population, baseline scenario", "Población proyectada para 2046 en el escenario base.", "Entero", "personas", 8190000, "No", "≥ 0",
     "«8,19 millones» convertido a personas (× 1.000.000).", "¿Cuánta población tendrá Hong Kong en 2046? (Wiki 1.1)"),
    ("F02", P, "poblacion_proyectada_2046_alta", "Projected population, high scenario", "Población proyectada para 2046 en el escenario alto.", "Entero", "personas", 8960000, "No", "≥ 0",
     "«8,96 M» convertido a personas.", "¿Cuál es el techo plausible de población en 2046? (Wiki 1.1)"),
    ("F02", P, "poblacion_proyectada_2046_baja", "Projected population, low scenario", "Población proyectada para 2046 en el escenario bajo.", "Entero", "personas", 7760000, "No", "≥ 0",
     "«7,76 M» convertido a personas.", "¿Cuál es el piso plausible de población en 2046? (Wiki 1.1)"),
    ("F02", P, "entrada_neta_proyectada_2022_2046", "Net movement (inflow)", "Personas que entrarían en neto a Hong Kong en el periodo de proyección.", "Entero", "personas", 1520000, "No", "≥ 0",
     "«1,52 millones» convertido a personas.", "¿Cuánto depende el crecimiento de la migración? (Wiki 1.1, 1.9)"),
    ("F02", P, "entrada_one_way_permit_2022_2046", "One-way Permit holders", "Parte de la entrada neta que llega con One-way Permit desde China continental.", "Entero", "personas", 890000, "No", "≥ 0",
     "«0,89 M» convertido a personas.", "¿Cuánto aporta la migración desde China continental? (Wiki 1.9)"),
    ("F02", P, "entrada_empleadas_domesticas_2022_2046", "Foreign domestic helpers", "Parte de la entrada neta que corresponde a empleadas domésticas extranjeras.", "Entero", "personas", 240000, "No", "≥ 0",
     "«0,24 M» convertido a personas.", "¿Cuánto pesan las empleadas domésticas en el crecimiento? (Wiki 1.9)"),
    ("F02", P, "poblacion_trabajadora_pico_2038", "Labour force (projected peak)", "Población trabajadora en su punto máximo proyectado (2038).", "Entero", "personas", 3660000, "No", "≥ 0",
     "«3,66 M» convertido a personas.", "¿Cuándo empieza a caer la fuerza laboral? (Wiki 1.1)"),
    ("F02", P, "poblacion_trabajadora_2046", "Labour force (projected 2046)", "Población trabajadora proyectada para 2046.", "Entero", "personas", 3580000, "No", "≥ 0",
     "«3,58 M» convertido a personas.", "¿Cuánta fuerza laboral habrá en 2046? (Wiki 1.1)"),
    ("F02", P, "poblacion_65_mas_2021", "Population aged 65 and over", "Personas de 65 años o más en 2021 (base de la proyección).", "Entero", "personas", 1450000, "No", "≥ 0",
     "«1,45 M» convertido a personas.", "¿Cuántas personas mayores hay? (Wiki 1.3)"),
    ("F02", P, "pct_65_mas_2021", "Proportion aged 65 and over", "Porcentaje de la población con 65 años o más en 2021.", "Decimal", "%", 20.5, "No", "0 a 100",
     "Coma decimal cambiada a punto. Base de proyección: excluye empleadas domésticas extranjeras.", "¿Qué tan envejecida está la ciudad? (Wiki 1.3)"),
    ("F02", P, "poblacion_65_mas_2046", "Population aged 65 and over (projected)", "Personas de 65 años o más proyectadas para 2046.", "Entero", "personas", 2740000, "No", "≥ 0",
     "«2,74 M» convertido a personas.", "¿Cuánto crecerá la población mayor? (Wiki 1.3)"),
    ("F02", P, "pct_65_mas_2046", "Proportion aged 65 and over (projected)", "Porcentaje proyectado de población con 65+ en 2046.", "Decimal", "%", 36, "No", "0 a 100",
     "Se copió el porcentaje como número.", "¿Qué presión tendrá el sistema de cuidado en 2046? (Wiki 1.3, 1.8)"),
    ("F03", P, "pct_15mas_secundaria_o_mas_2011", "Educational attainment: secondary and above", "Porcentaje de la población de 15+ con secundaria o más en 2011.", "Decimal", "%", 77.3, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Cómo ha mejorado el nivel educativo? (Wiki 1.6)"),
    ("F03", P, "pct_15mas_secundaria_o_mas_2021", "Educational attainment: secondary and above", "Porcentaje de la población de 15+ con secundaria o más en 2021.", "Decimal", "%", 81.6, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Qué nivel educativo tiene la población? (Wiki 1.6)"),
    ("F03", P, "pct_15mas_educacion_superior_2011", "Educational attainment: post-secondary", "Porcentaje de la población de 15+ con educación superior en 2011.", "Decimal", "%", 27.3, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Cuánto creció la educación superior en una década? (Wiki 1.6)"),
    ("F03", P, "pct_15mas_educacion_superior_2021", "Educational attainment: post-secondary", "Porcentaje de la población de 15+ con educación superior en 2021.", "Decimal", "%", 34.6, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Qué tanto capital humano tiene la ciudad? (Wiki 1.6)"),
    ("F03", P, "tasa_escolarizacion_3_5_2021", "School attendance rate, aged 3–5", "Porcentaje de niños de 3 a 5 años que asisten a la escuela en 2021.", "Decimal", "%", 88.4, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Qué tan extendido está el acceso a la educación inicial? (Wiki 1.6)"),
    ("F04", P, "tasa_bruta_natalidad_2025", "Crude birth rate", "Nacimientos por cada 1.000 habitantes en 2025 (provisional).", "Decimal", "por 1.000 hab.", 4.2, "No", "≥ 0",
     "Se quitó el «~»; coma decimal cambiada a punto. Reemplaza la TFR de agregadores (F17, descartada).", "¿Qué tan baja es la natalidad hoy? (Wiki 1.10)"),
    ("F04", P, "tasa_bruta_natalidad_2011", "Crude birth rate", "Nacimientos por cada 1.000 habitantes en 2011.", "Decimal", "por 1.000 hab.", 13.5, "No", "≥ 0",
     "Coma decimal cambiada a punto.", "¿Cuánto cayó la natalidad desde 2011? (Wiki 1.10)"),
    ("F04", P, "tasa_bruta_natalidad_1981", "Crude birth rate", "Nacimientos por cada 1.000 habitantes en 1981.", "Decimal", "por 1.000 hab.", 16.8, "No", "≥ 0",
     "Coma decimal cambiada a punto.", "¿Cómo ha cambiado la natalidad en 40 años? (Wiki 1.10)"),
    # ---------------- Índice 2 — Vivienda ----------------
    ("F05", V, "precio_m2_apto_menor_40m2_oct2025", "Average price, Class A (<40 m²)", "Precio promedio por m² de apartamentos de menos de 40 m² en octubre de 2025.", "Entero", "HK$/m²", 131305, "No", "≥ 0",
     "Se quitaron «HK$» y separadores de miles.", "¿Cuánto cuesta comprar vivienda? (Wiki 2.1)"),
    ("F05", V, "precio_m2_rango_min", "Average prices by class (lower bound)", "Límite inferior del rango de precio por m² según tamaño y zona.", "Entero", "HK$/m²", 130000, "No", "≥ 0",
     "Se separó el rango «HK$130.000–190.000/m²» en dos campos.", "¿Cuál es el piso del precio de la vivienda? (Wiki 2.1)"),
    ("F05", V, "precio_m2_rango_max", "Average prices by class (upper bound)", "Límite superior del rango de precio por m² según tamaño y zona.", "Entero", "HK$/m²", 190000, "No", "≥ 0",
     "Se separó el rango en dos campos.", "¿Cuál es el techo del precio de la vivienda? (Wiki 2.1)"),
    ("F05", V, "var_precio_h1_2026_pct", "Price index change", "Variación del índice de precios residenciales en el primer semestre de 2026.", "Decimal", "%", 7.9, "No", "−100 a 100",
     "Coma decimal cambiada a punto.", "¿Se está recuperando el precio de la vivienda? (Wiki 2.1)"),
    ("F05", V, "caida_precio_desde_pico_sep2021_pct", "Price index change from peak", "Caída del precio desde el máximo de septiembre de 2021.", "Decimal", "%", -24, "No", "−100 a 0",
     "Se guardó con signo negativo porque es una caída; se quitó el «~».", "¿Cuánto bajó la vivienda desde su pico? (Wiki 2.1)"),
    ("F05", V, "var_indice_rentas_h1_2026_pct", "Rental index change", "Variación del índice de rentas residenciales en el primer semestre de 2026.", "Decimal", "%", 2.6, "No", "−100 a 100",
     "Se quitó el «~»; coma decimal cambiada a punto.", "¿Está subiendo el arriendo? (Wiki 2.2)"),
    ("F05", V, "prevision_rentas_2026", "Rental forecast", "Rango de variación esperada de las rentas para todo 2026.", "Texto", "%", "+5 a +8", "Sí", "texto «+a a +b»",
     "Se copió como texto porque es un rango de previsión, no un dato medido.", "¿Qué se espera del arriendo este año? (Wiki 2.2)"),
    ("F05", V, "vacancia_residencial_pct_2025", "Vacancy rate", "Porcentaje del stock residencial privado vacío a cierre de 2025.", "Decimal", "%", 4.3, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Hay escasez de vivienda o falta de asequibilidad? (Wiki 2.9)"),
    ("F05", V, "viviendas_vacantes_2025", "Vacant units", "Viviendas privadas vacías a cierre de 2025.", "Entero", "unidades", 56080, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Cuántas viviendas están vacías? (Wiki 2.9)"),
    ("F05", V, "vacantes_sin_certificado_2025", "Units without Certificate of Compliance", "Viviendas vacantes que aún no tenían Certificado de Cumplimiento/Cesión.", "Entero", "unidades", 7120, "No", "≥ 0",
     "Se quitó el «~».", "¿Cuántas vacantes son obra recién terminada? (Wiki 2.9)"),
    ("F05", V, "oferta_privada_prevista_2026", "Forecast completions", "Viviendas privadas que se prevé terminar en 2026.", "Entero", "unidades", 16980, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Cuánta vivienda nueva llega al mercado? (Wiki 2.6)"),
    ("F05", V, "oferta_privada_prevista_2027", "Forecast completions", "Viviendas privadas que se prevé terminar en 2027.", "Entero", "unidades", 15360, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿La oferta va al alza o a la baja? (Wiki 2.6)"),
    ("F05", V, "viviendas_privadas_terminadas_2024", "Completions", "Viviendas privadas terminadas en 2024 (máximo en 20 años).", "Entero", "unidades", 24261, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Cuánta vivienda se construyó en el último pico? (Wiki 2.6)"),
    ("F05", V, "stock_viviendas_privadas_2025", "Private domestic stock", "Total aproximado de viviendas privadas a fin de 2025.", "Entero", "unidades", 1310000, "No", "≥ 0",
     "«~1,31 M» convertido a unidades.", "¿Qué tamaño tiene el parque de vivienda privada? (Wiki 2.6)"),
    ("F06", V, "pct_vivienda_publica_alquiler", "Public rental housing (PRH)", "Porcentaje de la población que vive en vivienda pública de alquiler.", "Decimal", "%", 29, "No", "0 a 100",
     "Se quitó el «~».", "¿Cuánta gente depende de vivienda pública de alquiler? (Wiki 2.4)"),
    ("F06", V, "pct_vivienda_subsidiada_propiedad", "Subsidised sale flats", "Porcentaje de la población en vivienda subsidiada en propiedad.", "Decimal", "%", 15.8, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Qué peso tiene la vivienda subsidiada? (Wiki 2.4)"),
    ("F06", V, "pct_vivienda_privada_permanente", "Private permanent housing", "Porcentaje de la población en vivienda privada permanente.", "Decimal", "%", 53.2, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Qué peso tiene el mercado privado? (Wiki 2.4)"),
    ("F06", V, "pct_vivienda_no_domestica", "Non-domestic housing", "Porcentaje de la población en alojamientos no domésticos.", "Decimal", "%", 1.2, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Cuánta gente vive fuera de hogares domésticos? (Wiki 2.4)"),
    ("F06", V, "pct_vivienda_temporal", "Temporary housing", "Porcentaje de la población en vivienda temporal.", "Decimal", "%", 0.7, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Persiste la vivienda precaria? (Wiki 2.4, 2.7)"),
    ("F06", V, "pct_poblacion_vivienda_publica_o_subsidiada", "Public housing (total)", "Porcentaje de la población que depende de vivienda pública o subsidiada.", "Decimal", "%", 45, "No", "0 a 100",
     "Se quitó el «~».", "¿Qué tanto depende la ciudad del Estado para vivir? (Wiki 2.4)"),
    ("F06", V, "renta_media_prh_mar2025", "Average monthly rent of PRH", "Renta mensual promedio de la vivienda pública de alquiler (marzo de 2025).", "Entero", "HK$ por mes", 2523, "No", "≥ 0",
     "Se quitaron «HK$», el «~» y separadores.", "¿Cuánto paga un inquilino de vivienda pública? (Wiki 2.4)"),
    ("F06", V, "solicitudes_prh_generales_mar2025", "General applications for PRH", "Solicitudes generales en lista de espera de vivienda pública (marzo de 2025).", "Entero", "solicitudes", 116400, "No", "≥ 0",
     "Se quitó el «~» y separadores.", "¿Cuánta gente espera vivienda pública? (Wiki 2.5)"),
    ("F06", V, "tiempo_espera_compuesto_anios_mar2026", "Composite Waiting Time (CWT)", "Tiempo de espera compuesto para obtener vivienda pública (marzo de 2026).", "Decimal", "años", 4.7, "No", "≥ 0",
     "Coma decimal cambiada a punto.", "¿Cuánto tarda una familia en recibir vivienda pública? (Wiki 2.5)"),
    ("F06", V, "tiempo_espera_prh_tradicional_anios", "Waiting time, traditional PRH", "Espera para vivienda pública tradicional estándar.", "Decimal", "años", 5.6, "No", "≥ 0",
     "Se quitó el «~»; coma decimal cambiada a punto.", "¿Mejora la espera para la vivienda pública tradicional? (Wiki 2.5)"),
    ("F06", V, "light_public_housing_unidades_1t2026", "Light Public Housing units", "Unidades de Light Public Housing disponibles al primer trimestre de 2026.", "Entero", "unidades", 9650, "No", "≥ 0",
     "Se quitó el «~» y separadores.", "¿Cuánto avanza la vivienda pública transitoria? (Wiki 2.5)"),
    ("F06", V, "light_public_housing_meta_2027_28", "Light Public Housing target", "Meta de unidades de Light Public Housing para 2027/28.", "Entero", "unidades", 30000, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Qué tan lejos está la meta de vivienda transitoria? (Wiki 2.5)"),
    ("F06", V, "oferta_publica_proyectada_5_anios", "Public housing supply (5 years)", "Viviendas públicas que se proyecta entregar en los próximos 5 años.", "Entero", "unidades", 196000, "No", "≥ 0",
     "Se quitó el «~» y separadores.", "¿Cuánta vivienda pública llegará? (Wiki 2.5, 2.6)"),
    # ---------------- Índice 3 — Finanzas públicas ----------------
    ("F08", FI, "ingreso_total_2026_27", "Total government revenue", "Ingreso total del Gobierno estimado para el año fiscal 2026-27.", "Entero", "HK$ millones", 765200, "No", "≥ 0",
     "Se expresó en millones de HK$ sin separadores.", "¿Con cuánto dinero cuenta el Gobierno? (Wiki 3.1)"),
    ("F08", FI, "gasto_pct_pib_2026_27", "Expenditure as % of GDP", "Gasto público como porcentaje del PIB.", "Decimal", "% del PIB", 24.2, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Qué tan grande es el Estado en la economía? (Wiki 3.1)"),
    ("F08", FI, "impuestos_ganancias_y_salarios_2026_27", "Profits tax and salaries tax", "Recaudación conjunta de impuestos sobre ganancias y salarios.", "Entero", "HK$ millones", 321200, "No", "≥ 0",
     "Se quitó el «~»; expresado en millones.", "¿Cuál es la principal fuente de ingreso fiscal? (Wiki 3.1, 3.2)"),
    ("F08", FI, "derechos_timbre_2026_27", "Stamp duties", "Recaudación estimada por derechos de timbre.", "Entero", "HK$ millones", 101000, "No", "≥ 0",
     "Se quitó el «~»; expresado en millones.", "¿Cuánto depende el fisco del mercado bursátil e inmobiliario? (Wiki 3.3)"),
    ("F08", FI, "ingreso_suelo_2026_27", "Land premium", "Ingreso estimado por venta/arriendo de suelo.", "Entero", "HK$ millones", 18000, "No", "≥ 0",
     "Expresado en millones.", "¿Qué tan volátil es el ingreso por suelo? (Wiki 3.4)"),
    ("F08", FI, "gasto_total_2026_27", "Total government expenditure", "Gasto total del Gobierno estimado para 2026-27.", "Entero", "HK$ millones", 843400, "No", "≥ 0",
     "Expresado en millones sin separadores.", "¿Cuánto gasta el Gobierno? (Wiki 3.5)"),
    ("F08", FI, "var_gasto_total_pct", "Change in total expenditure", "Variación del gasto total frente al año anterior.", "Decimal", "%", 6.9, "No", "−100 a 100",
     "Coma decimal cambiada a punto.", "¿El gasto público se está expandiendo? (Wiki 3.5)"),
    ("F08", FI, "gasto_recurrente_2026_27", "Recurrent expenditure", "Gasto recurrente estimado.", "Entero", "HK$ millones", 599700, "No", "≥ 0",
     "Expresado en millones.", "¿Cuánto del gasto es permanente? (Wiki 3.5)"),
    ("F08", FI, "gasto_salud_bienestar_educacion_2026_27", "Recurrent expenditure on health, welfare and education", "Gasto en salud, bienestar y educación (~60 % del recurrente).", "Entero", "HK$ millones", 357100, "No", "≥ 0",
     "Expresado en millones.", "¿Cuánto se invierte en lo social? (Wiki 3.5)"),
    ("F08", FI, "deficit_capital_2026_27", "Capital account deficit", "Déficit de la cuenta de capital (obras como Northern Metropolis).", "Entero", "HK$ millones", -90100, "No", "cualquier entero",
     "Se guardó con signo negativo porque es déficit; se quitó el «~».", "¿Cuánto cuesta la inversión en infraestructura? (Wiki 3.5, 3.6)"),
    ("F08", FI, "superavit_consolidado_2025_26", "Consolidated surplus (revised)", "Resultado fiscal consolidado revisado del año 2025-26.", "Entero", "HK$ millones", 2900, "No", "cualquier entero",
     "Expresado en millones; positivo = superávit.", "¿Están sanas las cuentas públicas? (Wiki 3.6)"),
    ("F08", FI, "superavit_consolidado_2026_27", "Consolidated surplus (estimate, after bonds)", "Resultado fiscal consolidado estimado para 2026-27 tras emisión de bonos.", "Entero", "HK$ millones", 22100, "No", "cualquier entero",
     "Expresado en millones; positivo = superávit.", "¿El Gobierno cierra el año en positivo? (Wiki 3.6)"),
    ("F08", FI, "reservas_fiscales_mar2026", "Fiscal reserves", "Reservas fiscales al 31 de marzo de 2026.", "Entero", "HK$ millones", 657200, "No", "≥ 0",
     "Se quitó el «~»; expresado en millones.", "¿Qué colchón financiero tiene la ciudad? (Wiki 3.7)"),
    ("F08", FI, "reservas_meses_de_gasto", "Fiscal reserves in months of expenditure", "Meses de gasto público que cubrirían las reservas al cierre de 2026-27 (HK$679.300 M).", "Entero", "meses", 10, "No", "≥ 0",
     "Se quitó el «~».", "¿Cuánto aguantaría el Gobierno sin ingresos? (Wiki 3.7)"),
    ("F08", FI, "techo_deuda_bonos", "Borrowing ceiling", "Límite máximo de endeudamiento por bonos del Gobierno.", "Entero", "HK$ millones", 900000, "No", "≥ 0",
     "Expresado en millones.", "¿Cuánto se puede endeudar el Gobierno? (Wiki 3.8)"),
    ("F08", FI, "emision_bonos_2026_27", "Bond issuance", "Bonos que el Gobierno prevé emitir en 2026-27.", "Entero", "HK$ millones", 160000, "No", "≥ 0",
     "Se quitó el «~»; expresado en millones.", "¿Cuánta deuda nueva toma el Gobierno? (Wiki 3.8)"),
    ("F09", FI, "pct_pib_cuatro_industrias_2024", "Share of GDP of the four key industries", "Peso en el PIB de finanzas, comercio y logística, servicios profesionales y turismo.", "Decimal", "% del PIB", 58.2, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿De qué vive la economía de Hong Kong? (Wiki 3.1)"),
    ("F09", FI, "pct_empleo_cuatro_industrias_2024", "Share of employment of the four key industries", "Peso en el empleo total de las cuatro industrias clave.", "Decimal", "% del empleo", 42.3, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Cuánto empleo generan los sectores clave? (Wiki 3.1)"),
    # ---------------- Índice 4 — Seguridad ----------------
    ("F10", S, "delitos_totales_2025", "Overall crime", "Delitos totales registrados en 2025.", "Entero", "casos", 89137, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Qué tan segura es la ciudad? (Wiki 4.1)"),
    ("F10", S, "var_delitos_totales_2025_pct", "Change in overall crime", "Variación de los delitos totales frente a 2024.", "Decimal", "%", -5.9, "No", "−100 a 100",
     "Coma decimal cambiada a punto; negativo = baja.", "¿La criminalidad sube o baja? (Wiki 4.1)"),
    ("F10", S, "delitos_violentos_2025", "Violent crime", "Delitos violentos registrados en 2025.", "Entero", "casos", 8823, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Cuánto delito violento hay? (Wiki 4.2)"),
    ("F10", S, "var_delitos_violentos_2025_pct", "Change in violent crime", "Variación de los delitos violentos frente a 2024.", "Decimal", "%", -15.9, "No", "−100 a 100",
     "Coma decimal cambiada a punto; negativo = baja.", "¿El delito violento está bajando? (Wiki 4.2)"),
    ("F10", S, "var_robo_con_violencia_2025_pct", "Change in robbery", "Variación del robo con violencia frente a 2024.", "Decimal", "%", -26.7, "No", "−100 a 100",
     "Coma decimal cambiada a punto.", "¿Qué delitos callejeros están en mínimos? (Wiki 4.2)"),
    ("F10", S, "var_hurto_con_allanamiento_2025_pct", "Change in burglary", "Variación del hurto con allanamiento frente a 2024.", "Decimal", "%", -33.1, "No", "−100 a 100",
     "Coma decimal cambiada a punto.", "¿Qué delitos contra la propiedad están en mínimos? (Wiki 4.2)"),
    ("F10", S, "casos_fraude_2025", "Deception cases", "Casos de fraude y estafa registrados en 2025.", "Entero", "casos", 43212, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Qué tan grave es el problema de las estafas? (Wiki 4.4)"),
    ("F10", S, "pct_fraude_sobre_delitos_2025", "Deception as share of overall crime", "Porcentaje de los delitos que son fraude o estafa.", "Decimal", "%", 48.5, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Está cambiando el tipo de delito en la ciudad? (Wiki 4.1, 4.4)"),
    ("F10", S, "perdidas_fraude_2025", "Losses from deception", "Pérdidas económicas por fraude en 2025.", "Entero", "HK$ millones", 8100, "No", "≥ 0",
     "«HK$8.100 M» expresado en millones sin separadores.", "¿Cuánto dinero se pierde por estafas? (Wiki 4.4)"),
    ("F10", S, "var_perdidas_fraude_2025_pct", "Change in losses from deception", "Variación de las pérdidas por fraude frente a 2024.", "Decimal", "%", -11.3, "No", "−100 a 100",
     "Coma decimal cambiada a punto.", "¿Las estafas están costando menos? (Wiki 4.4)"),
    ("F10", S, "casos_fraude_compras_online_2025", "Online shopping fraud", "Casos de fraude en compras por internet (el tipo más común).", "Entero", "casos", 12505, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Cuál es la estafa más frecuente? (Wiki 4.4)"),
    ("F10", S, "delitos_triadas_2025", "Triad-related crime", "Delitos relacionados con tríadas en 2025.", "Entero", "casos", 1944, "No", "≥ 0",
     "Se quitaron separadores de miles.", "¿Cómo evoluciona el crimen organizado? (Wiki 4.5)"),
    ("F10", S, "var_delitos_triadas_2025_pct", "Change in triad-related crime", "Variación de los delitos de tríadas frente a 2024.", "Decimal", "%", -16.4, "No", "−100 a 100",
     "Coma decimal cambiada a punto.", "¿El crimen organizado está bajando? (Wiki 4.5)"),
    ("F10", S, "tasa_homicidios_habitual_100k", "Homicide rate", "Tasa habitual de homicidios por cada 100.000 habitantes (sin eventos únicos).", "Decimal", "por 100.000 hab.", 0.4, "No", "≥ 0",
     "Se quitó el «~». No incluye el incendio de Wang Fuk Court de 2025, reclasificado como homicidio involuntario.", "¿Qué tan común es el homicidio? (Wiki 4.3)"),
    ("F11", S, "posicion_cpi_2025", "CPI 2025 rank", "Puesto de Hong Kong en el Índice de Percepción de la Corrupción 2025.", "Entero", "puesto", 12, "No", "1 a 180+",
     "Se tomó el número del puesto («12º»).", "¿Qué tan corrupta se percibe la ciudad? (Wiki 4.6)"),
    ("F11", S, "puntaje_cpi_2025", "CPI 2025 score", "Puntaje de Hong Kong en el Índice de Percepción de la Corrupción 2025.", "Entero", "puntos (0–100)", 76, "No", "0 a 100",
     "Copiado como entero.", "¿Cómo se compara la corrupción con otras ciudades? (Wiki 4.6)"),
    ("F11", S, "posicion_cpi_2024", "CPI 2024 rank", "Puesto de Hong Kong en el Índice de Percepción de la Corrupción 2024.", "Entero", "puesto", 17, "No", "1 a 180+",
     "Se tomó el número del puesto («17º»).", "¿Mejoró la percepción de corrupción? (Wiki 4.6)"),
    # ---------------- Índice 5 — Turismo ----------------
    ("F12", T, "llegadas_visitantes_2025", "Visitor arrivals", "Visitantes que llegaron a Hong Kong en 2025.", "Entero", "visitantes", 49900000, "No", "≥ 0",
     "«49,9 millones» convertido a visitantes.", "¿Cuántos turistas recibe la ciudad? (Wiki 5.2)"),
    ("F12", T, "var_llegadas_2025_pct", "Change in visitor arrivals", "Variación de las llegadas frente a 2024.", "Decimal", "%", 12, "No", "−100 a 500",
     "Se copió el porcentaje como número.", "¿Se está recuperando el turismo? (Wiki 5.2)"),
    ("F12", T, "llegadas_visitantes_pico_2018", "Visitor arrivals (2018)", "Llegadas en 2018, el máximo histórico de referencia.", "Entero", "visitantes", 65300000, "No", "≥ 0",
     "«65,3 M» convertido a visitantes.", "¿Qué tan lejos está el turismo de su pico? (Wiki 5.2)"),
    ("F12", T, "llegadas_visitantes_1t2026", "Visitor arrivals (Q1 2026)", "Llegadas en el primer trimestre de 2026 (récord).", "Entero", "visitantes", 14310000, "No", "≥ 0",
     "«~14,31 M» convertido a visitantes.", "¿Cómo va el turismo este año? (Wiki 5.2)"),
    ("F12", T, "var_llegadas_1t2026_pct", "Change in visitor arrivals (Q1 2026)", "Variación de las llegadas del primer trimestre de 2026 frente a 2025.", "Decimal", "%", 17, "No", "−100 a 500",
     "Se copió el porcentaje como número.", "¿El turismo sigue acelerando? (Wiki 5.2)"),
    ("F12", T, "llegadas_continentales_2025", "Mainland visitors", "Visitantes de China continental en 2025.", "Entero", "visitantes", 37800000, "No", "≥ 0",
     "«37,8 M» convertido a visitantes.", "¿Qué tanto depende el turismo de China continental? (Wiki 5.3)"),
    ("F12", T, "llegadas_no_continentales_2025", "Non-mainland visitors", "Visitantes de otros mercados en 2025.", "Entero", "visitantes", 12100000, "No", "≥ 0",
     "«12,1 M» convertido a visitantes.", "¿Cuánto turismo internacional llega? (Wiki 5.3)"),
    ("F12", T, "pct_pib_turismo_2024", "Tourism share of GDP", "Peso del turismo en el PIB (medida oficial).", "Decimal", "% del PIB", 2.8, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Qué tan importante es el turismo para la economía? (Wiki 5.1)"),
    ("F12", T, "empleos_turismo_2024", "Tourism employment", "Empleos generados por el turismo en 2024.", "Entero", "empleos", 159700, "No", "≥ 0",
     "Se quitó el «~» y separadores.", "¿Cuánto empleo genera el turismo? (Wiki 5.1)"),
    ("F12", T, "pct_empleo_turismo_2024", "Tourism share of employment", "Peso del turismo en el empleo total.", "Decimal", "% del empleo", 4.3, "No", "0 a 100",
     "Coma decimal cambiada a punto.", "¿Qué parte del empleo depende del turismo? (Wiki 5.1)"),
    ("F12", T, "gasto_turistico_receptor_2024", "Inbound tourism expenditure", "Gasto total de los visitantes en 2024.", "Entero", "HK$ millones", 195000, "No", "≥ 0",
     "«HK$195.000 M» expresado en millones.", "¿Cuánto gastan los turistas? (Wiki 5.4)"),
    ("F12", T, "var_gasto_turistico_2024_pct", "Change in inbound tourism expenditure", "Variación del gasto turístico frente a 2023.", "Decimal", "%", 9.6, "No", "−100 a 500",
     "Coma decimal cambiada a punto.", "¿El gasto turístico está creciendo? (Wiki 5.4)"),
    ("F12", T, "estancia_media_noches_2025", "Average length of stay", "Noches promedio que se queda un visitante con pernoctación.", "Decimal", "noches", 3.1, "No", "≥ 0",
     "Se quitó el «~»; coma decimal cambiada a punto.", "¿Cuánto tiempo se quedan los turistas? (Wiki 5.5)"),
    ("F12", T, "ocupacion_hotelera_2025_pct", "Hotel room occupancy rate", "Porcentaje de ocupación de las habitaciones de hotel en 2025.", "Decimal", "%", 87, "No", "0 a 100",
     "Se copió el porcentaje como número.", "¿Alcanza la oferta hotelera? (Wiki 5.6)"),
    # ---------------- Índice 6 — Suelo ----------------
    ("F13", SU, "n_country_parks", "Number of country parks", "Número de parques naturales (country parks) designados.", "Entero", "parques", 25, "No", "≥ 0",
     "Copiado como entero.", "¿Cuánto territorio está protegido? (Wiki 6.6)"),
    ("F13", SU, "n_special_areas", "Number of special areas", "Número de áreas especiales de conservación.", "Entero", "áreas", 22, "No", "≥ 0",
     "Copiado como entero.", "¿Cuántas áreas especiales hay? (Wiki 6.6)"),
    ("F13", SU, "area_protegida_km2", "Total area of country parks and special areas", "Superficie total de parques y áreas especiales.", "Entero", "km²", 442, "No", "≥ 0",
     "Se quitó el «~».", "¿Cuánto suelo está protegido? (Wiki 6.6)"),
    ("F13", SU, "pct_territorio_protegido", "Share of land protected", "Porcentaje del territorio bajo protección.", "Decimal", "%", 40, "No", "0 a 100",
     "Se quitó el «~».", "¿Qué proporción de la ciudad es naturaleza protegida? (Wiki 6.6)"),
    ("F13", PQ, "puesto", "(orden de la tabla)", "Posición del parque en el ranking por tamaño.", "Entero", "puesto", 1, "No", "1 a 6",
     "Se tomó la columna «#» de la tabla de la wiki (6 parques más grandes).", "¿Cuáles son los parques naturales más grandes? (Wiki 5.7, 6.6)"),
    ("F13", PQ, "nombre_parque", "Country Park", "Nombre del parque natural.", "Texto", "—", "Lantau South", "No", "Lantau South, Tai Lam, Plover Cove, Sai Kung East, Pat Sin Leng, Sai Kung West",
     "Copiado tal cual en inglés.", "¿Qué parques visitar o proteger? (Wiki 5.7)"),
    ("F13", PQ, "area_ha", "Area (ha)", "Superficie del parque natural.", "Entero", "hectáreas", 5646, "No", "3.000 a 5.646",
     "Se quitaron separadores de miles.", "¿Qué tan grande es cada parque? (Wiki 5.7, 6.6)"),
    ("F14", SU, "pct_suministro_dongjiang_min", "Share of supply from Dongjiang (lower bound)", "Límite inferior del porcentaje del agua que viene del río Dongjiang.", "Decimal", "%", 70, "No", "0 a 100",
     "Se separó el rango «~70–80 %» en dos campos.", "¿Qué tan dependiente es la ciudad del agua de China continental? (Wiki 6.7)"),
    ("F14", SU, "pct_suministro_dongjiang_max", "Share of supply from Dongjiang (upper bound)", "Límite superior del porcentaje del agua que viene del río Dongjiang.", "Decimal", "%", 80, "No", "0 a 100",
     "Se separó el rango en dos campos.", "¿Qué tan dependiente es la ciudad del agua de China continental? (Wiki 6.7)"),
    ("F14", SU, "agua_importada_dongjiang_2024_25", "Dongjiang water imported", "Volumen de agua importada del Dongjiang en 2024/25.", "Entero", "millones de m³", 816, "No", "≥ 0",
     "«816 M m³» expresado en millones de m³.", "¿Cuánta agua se importa? (Wiki 6.7)"),
    ("F14", SU, "pct_consumo_dongjiang_2024", "Water consumption mix: Dongjiang", "Parte del consumo de agua de 2024 cubierta por el Dongjiang.", "Decimal", "%", 60, "No", "0 a 100",
     "Se quitó el «~».", "¿De dónde sale el agua que se consume? (Wiki 6.7)"),
    ("F14", SU, "pct_consumo_lluvia_local_2024", "Water consumption mix: local yield", "Parte del consumo cubierta por lluvia local.", "Decimal", "%", 17, "No", "0 a 100",
     "Se quitó el «~».", "¿Cuánto aporta la lluvia local? (Wiki 6.7)"),
    ("F14", SU, "pct_consumo_agua_mar_2024", "Water consumption mix: seawater for flushing", "Parte del consumo cubierta con agua de mar (inodoros).", "Decimal", "%", 22, "No", "0 a 100",
     "Se quitó el «~».", "¿Cuánto se ahorra usando agua de mar? (Wiki 6.7)"),
    ("F14", SU, "pct_consumo_desalacion_2024", "Water consumption mix: desalination", "Parte del consumo cubierta por agua desalada.", "Decimal", "%", 1, "No", "0 a 100",
     "Se quitó el «~».", "¿Qué tanto aporta la desalación hoy? (Wiki 6.7)"),
    ("F14", SU, "capacidad_embalses", "Total storage capacity of reservoirs", "Capacidad total de los embalses.", "Entero", "millones de m³", 586, "No", "≥ 0",
     "«586 M m³» expresado en millones de m³.", "¿Cuánta agua puede almacenar la ciudad? (Wiki 6.7)"),
    ("F15", SU, "especies_evaluadas", "Species assessed", "Especies evaluadas en el informe de biodiversidad 2025.", "Entero", "especies", 886, "No", "≥ 0",
     "Copiado como entero.", "¿Cuánta biodiversidad se ha evaluado? (Wiki 6.11)"),
    ("F15", SU, "especies_extintas_localmente", "Locally extinct species", "Especies evaluadas que ya se extinguieron localmente.", "Entero", "especies", 21, "No", "≥ 0",
     "Copiado como entero.", "¿Cuánta biodiversidad se ha perdido? (Wiki 6.11)"),
    ("F15", SU, "pct_especies_en_riesgo", "Share of species at risk", "Porcentaje de especies evaluadas que están en riesgo.", "Decimal", "%", 26, "No", "0 a 100",
     "Se copió el porcentaje como número.", "¿Qué tan amenazada está la naturaleza local? (Wiki 6.11)"),
]


def main(entrada, salida):
    wb = openpyxl.load_workbook(entrada)
    ws = wb["2. Diccionario"]
    n = len(CAMPOS)
    fin = 4 + n
    modelo = [ws.cell(row=5, column=c) for c in range(1, 14)]
    for r in range(5, fin + 1):
        for c, m in enumerate(modelo, 1):
            cel = ws.cell(row=r, column=c)
            if r > 44:
                cel.fill, cel.font, cel.border, cel.alignment = copy(m.fill), copy(m.font), copy(m.border), copy(m.alignment)
    for i, (fid, archivo, campo, orig, sig, tipo, unidad, ej, vacio, rango, trans, preg) in enumerate(CAMPOS, 5):
        fila = [fid, archivo, campo, orig, sig, tipo, unidad, ej, vacio, rango, "Pública", trans, preg]
        for j, v in enumerate(fila, 1):
            ws.cell(row=i, column=j, value=v)
    if fin > 44:
        for dv in ws.data_validations.dataValidation:
            dv.sqref = MultiCellRange(re.sub(r"([A-Z]+)5:([A-Z]+)44", lambda m: f"{m.group(1)}5:{m.group(2)}{fin}", str(dv.sqref)))
        res = wb["Resumen"]
        for row in res.iter_rows(min_col=2, max_col=2):
            for c in row:
                if isinstance(c.value, str) and "'2. Diccionario'" in c.value:
                    c.value = re.sub(r"('2\. Diccionario'!\$?[A-Z]+\$?5:\$?[A-Z]+\$?)44", rf"\g<1>{fin}", c.value)
    wb.save(salida)
    print(f"{n} campos escritos (filas 5 a {fin})")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
