"""Llena la plantilla del Taller de Datos (Fuentes, Diccionario, Bitácora IA) con el lago de Hong Kong.

Uso: python pipeline/llenar_plantilla.py <plantilla.xlsx> <salida.xlsx>
"""
from __future__ import annotations

import json
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
FECHA = "2026-10-06"
REPO = "https://github.com/ManuelArenas1/Hong-Kong-gmdigital"
DATAGOV = "Terms and Conditions of Use of DATA.GOV.HK (reutilización comercial y no comercial con atribución)"

# ---------------------------------------------------------------------------
# 1. Fuentes (integradas y no integradas)
# ---------------------------------------------------------------------------
FUENTES = [
    # id, nombre, entidad, tipo_pub, url, formato, obtencion, licencia, personas, cobertura, periodo, frecuencia, estado, como, obs
    ("F01", "Límites de circunscripciones del Consejo de Distrito (DCCA) 2019", "Electoral Affairs Commission (EAC); GeoJSON publicado por hkdce en GitHub",
     "Entidad pública", "https://raw.githubusercontent.com/hkdce/dcca-boundaries/master/converted-geojson/DCCA_2019.geojson", "GeoJSON",
     "Descarga directa", "no declara", "No", "18 distritos / 452 DCCA", "2019", "Única vez", "Integrada al lago",
     "La IA la ubicó por búsqueda web tras fallar otras opciones; descargada con curl",
     "El repositorio espejo no tiene archivo LICENSE (404). Los polígonos incluyen mar: la densidad usa área terrestre oficial (F04)."),
    ("F02", "Censo de Población 2021 — estadísticas por distrito (DC_21C)", "Census and Statistics Department (C&SD), vía portal CSDI",
     "Entidad pública", "https://portal.csdi.gov.hk/server/services/common/censtatd_rcd_1635932856094_19982/MapServer/WFSServer?service=wfs&request=GetFeature&typenames=csdi:DC_21C&outputFormat=CSV",
     "CSV", "API sin llave", "open data for free re-use … subject to the Terms and Conditions of Use as published on DATA.GOV.HK",
     "Solo conteos por zona", "18 distritos", "2021 (jun–ago)", "Única vez", "Integrada al lago",
     "La IA encontró el servicio WFS en la ficha de data.gov.hk después de que el XLSX/ZIP oficial no se pudiera leer",
     "Validado: suma de 5 grupos de edad y de 5 tipos de vivienda = población total en los 18 distritos."),
    ("F03", "Estadísticas de población y hogares por distrito (Encuesta General de Hogares)", "Census and Statistics Department (C&SD), vía portal CSDI",
     "Entidad pública", "https://portal.csdi.gov.hk/server/services/common/censtatd_rcd_1635934545173_69201/MapServer/WFSServer?service=wfs&request=GetFeature&typenames=csdi:DC_GHS&outputFormat=CSV",
     "CSV", "API sin llave", DATAGOV, "Solo conteos por zona", "18 distritos", "2016 a 2025", "Anual", "Integrada al lago",
     "La IA la encontró en data.gov.hk al buscar la publicación B1130301",
     "Validado: población 2021 del panel ≈ censo (F02). La edad mediana viene redondeada a entero."),
    ("F04", "Districts of Hong Kong (tabla de área, población y densidad)", "Wikipedia (Wikimedia Foundation)", "Plataforma colaborativa",
     "https://en.wikipedia.org/wiki/Districts_of_Hong_Kong", "Página web", "Solo PDF o página", "CC BY-SA 4.0", "No", "18 distritos",
     "Censo 2021", "No se sabe", "Integrada al lago", "La IA la consultó para obtener el área terrestre oficial",
     "La población 2021 coincide exacto con F02 en los 18 distritos; se usa solo el área y el nombre en chino."),
    ("F05", "Rutas y paradas de transporte público (bus, minibús, MTR, tren ligero, ferry)", "hkbus/hk-bus-crawling (compila datos abiertos de operadores)",
     "Plataforma colaborativa", "https://raw.githubusercontent.com/hkbus/hk-bus-crawling/gh-pages/routeFareList.min.json", "JSON", "Descarga directa",
     "GNU GENERAL PUBLIC LICENSE Version 2", "No", "Todo Hong Kong (15.287 paradas)", "Vigente a 2026-10", "No se sabe", "Integrada al lago",
     "La IA llegó a ella revisando el paquete npm hk-bus-eta; descargada con curl",
     "Los datos de origen son de KMB, Citybus, NLB, GMB, MTR y ferris publicados en DATA.GOV.HK."),
    ("F06", "Hospitales públicos con urgencias (A&E) y coordenadas", "Paquete npm ane-hk (lista de la Hospital Authority)", "Plataforma colaborativa",
     "https://registry.npmjs.org/ane-hk", "Otro", "Descarga directa", "GPL-3.0-only", "No", "18 hospitales", "Vigente a 2026-10", "No se sabe",
     "Integrada al lago", "La IA buscó en npm paquetes con datos de Hong Kong",
     "Solo se extrajeron nombres y coordenadas (hechos), no código. Coordenadas no contrastadas con fuente oficial; falta el nuevo hospital de Kai Tak."),
    ("F07", "Información de estaciones meteorológicas", "Hong Kong Observatory (HKO)", "Entidad pública", "https://www.hko.gov.hk/en/cis/stn.htm",
     "Página web", "Solo PDF o página", "no declara", "No", "86 estaciones", "1884 a 2026", "No se sabe", "Integrada al lago",
     "La IA la encontró por búsqueda web", "Coordenadas en grados-minutos-segundos convertidas a decimal; contrastadas con F08 (24 de 26 a <1 m)."),
    ("F08", "Tiempo regional — humedad relativa del último minuto (formato espacial)", "Hong Kong Observatory, vía portal CSDI", "Entidad pública",
     "https://portal.csdi.gov.hk/server/services/common/hko_rcd_1634875711587_25731/MapServer/WFSServer?service=wfs&request=GetFeature&typenames=csdi:latest_1min_humidity&outputFormat=GEOJSON",
     "GeoJSON", "API sin llave", DATAGOV, "No", "26 estaciones", "Tiempo real", "Tiempo real", "Integrada al lago",
     "La IA la encontró en data.gov.hk", "Se usa para control de calidad de coordenadas; trae la URL del CSV en vivo de cada estación."),
    ("F09", "AQHI actual por estación de calidad del aire", "Environmental Protection Department (EPD)", "Entidad pública",
     "https://www.aqhi.gov.hk/epd/ddata/html/out/aqhi_ind_rss_Eng.xml", "Otro", "Descarga directa", DATAGOV, "No", "18 estaciones",
     "Captura 2026-10-06 07:30 HKT", "Tiempo real", "Integrada al lago", "La IA la encontró en data.gov.hk (actualización horaria)",
     "Formato XML/RSS. Captura puntual; las estaciones se asignan a distrito por nombre (sin coordenadas)."),
    ("F10", "PIB per cápita de Hong Kong (US$ corrientes)", "Banco Mundial, publicado por FRED (St. Louis Fed)", "Entidad pública",
     "https://fred.stlouisfed.org/data/PCAGDPHKA646NWDB", "Página web", "Solo PDF o página",
     "Data in this graph are copyrighted. Please review the copyright information in the series notes before sharing.", "No", "Hong Kong",
     "1960 a 2025", "Anual", "Integrada al lago", "Búsqueda web después de que la API del Banco Mundial no se pudo usar", "Se usa 2000–2025."),
    ("F11", "PIB per cápita de Hong Kong (US$ constantes)", "Banco Mundial, publicado por FRED (St. Louis Fed)", "Entidad pública",
     "https://fred.stlouisfed.org/data/NYGDPPCAPKDHKG", "Página web", "Solo PDF o página",
     "Data in this graph are copyrighted. Please review the copyright information in the series notes before sharing.", "No", "Hong Kong",
     "1961 a 2025", "Anual", "Integrada al lago", "Búsqueda web", "FRED la rotula como dólares de 2010, pero el valor de 2015 coincide con el corriente: base 2015."),
    ("F12", "Inflación, precios al consumidor (% anual)", "Banco Mundial, publicado por FRED (St. Louis Fed)", "Entidad pública",
     "https://fred.stlouisfed.org/data/FPCPITOTLZGHKG", "Página web", "Solo PDF o página",
     "Data in this graph are copyrighted. Please review the copyright information in the series notes before sharing.", "No", "Hong Kong",
     "2000 a 2025 (usado)", "Anual", "Integrada al lago", "Búsqueda web", ""),
    ("F13", "Situación de delitos totales y violentos (mensual)", "Hong Kong Police Force (HKPF)", "Entidad pública",
     "https://www.police.gov.hk/info/doc/crime_details_overall.csv", "CSV", "Descarga directa", DATAGOV, "Solo conteos por zona", "Hong Kong (sin desagregar)",
     "2022-01 a 2024-12", "Mensual", "Integrada al lago", "La IA la encontró en data.gov.hk",
     "La suma 2024 (94.948) difiere del total anual publicado en F14 (94.747). No trae distritos."),
    ("F14", "Police in Figures 2025", "Hong Kong Police Force (HKPF)", "Entidad pública",
     "https://www.police.gov.hk/info/doc/police_in_figure/2025/2025_en.html", "Página web", "Solo PDF o página", "no declara", "Solo conteos por zona",
     "Hong Kong", "2024 y 2025", "Anual", "Integrada al lago", "Búsqueda web de estadísticas de delito por distrito",
     "No trae desagregación por distrito; licencia no revisada en la página. Coincide con la wiki (89.137 delitos, −5,9 %)."),
    ("F15", "Tasa de desempleo de Hong Kong (desestacionalizada)", "Moody's Analytics (economy.com), con datos del C&SD", "Empresa privada",
     "https://www.economy.com/hong-kong/unemployment-rate", "Página web", "Solo PDF o página", "no declara", "No", "Hong Kong",
     "2026-05 a 2026-08", "Mensual", "Integrada al lago", "Búsqueda web",
     "Solo los 2 últimos datos. Reemplazar por la tabla oficial 210-06101 del C&SD cuando se pueda leer su API."),
    ("F16", "Titulares de prensa sobre Hong Kong por tema", "Varios medios (The Standard, SCMP, RTHK, TVB, info.gov.hk, otros)", "Medio de comunicación",
     "(32 URL, una por titular, en curated/escucha/noticias.csv)", "Página web", "Solo PDF o página", "no declara", "Texto libre sin revisar",
     "Hong Kong", "2025 a 2026-10", "Única vez", "Integrada al lago", "8 búsquedas web hechas por la IA (vivienda, comercio, transporte, etc.)",
     "Solo titular y URL; cada medio tiene sus propios derechos. Los titulares pueden nombrar personas: revisar antes de publicar."),
    ("F17", "Diccionario de barrios por distrito (hk-location.json)", "Paquete npm hkopendata 1.0.0", "Plataforma colaborativa",
     "https://registry.npmjs.org/hkopendata/-/hkopendata-1.0.0.tgz", "JSON", "Descarga directa", "MIT", "No", "18 distritos", "No se sabe",
     "Única vez", "Integrada al lago", "La IA buscó en npm", "Se usa para detectar distritos mencionados en los titulares."),
    # --- No integradas ---
    ("F18", "Censo 2021 — perfiles por distrito (DC_21C.xlsx / DC_21C.zip)", "Census and Statistics Department (C&SD)", "Entidad pública",
     "https://www.census2021.gov.hk/doc/DC_21C.xlsx", "XLSX", "Bloquea scripts o pide sesión", DATAGOV, "Solo conteos por zona", "18 distritos",
     "2021", "Única vez", "Candidata", "La IA la encontró en la página District Profiles del censo",
     "Es la fuente oficial ideal, pero el entorno de la IA no puede descargarla y la herramienta web no lee binarios. Mismos datos que F02."),
    ("F19", "API de tablas web del C&SD (get.php)", "Census and Statistics Department (C&SD)", "Entidad pública",
     "https://www.censtatd.gov.hk/api/get.php?id=130-06805&lang=en&full_series=1", "JSON", "API sin llave", DATAGOV, "Solo conteos por zona", "Hong Kong",
     "Varía por tabla", "Mensual", "Candidata", "La IA la propuso para series oficiales (desempleo, hogares)",
     "Respondió «Parameter is not defined»: exige un parámetro comprimido que arma la web del C&SD."),
    ("F20", "geoBoundaries — Hong Kong ADM1", "William & Mary geoLab", "Academia u ONG",
     "https://raw.githubusercontent.com/wmgeolab/geoBoundaries/main/releaseData/gbOpen/HKG/ADM1/geoBoundaries-HKG-ADM1.geojson", "GeoJSON",
     "Descarga directa", "no declara", "No", "Hong Kong", "—", "No se sabe", "No existe (la inventó la IA)",
     "La IA supuso la ruta del archivo", "Respondió 404: no hay ADM1 de HKG en esa ruta."),
    ("F21", "Highcharts Map Collection — mapa de distritos de Hong Kong", "Highsoft (paquete npm @highcharts/map-collection)", "Empresa privada",
     "https://registry.npmjs.org/@highcharts/map-collection", "JSON", "Descarga directa", "no declara", "No", "Hong Kong", "—", "No se sabe",
     "No existe (la inventó la IA)", "La IA supuso que el paquete traía countries/hk", "La versión 2.3.3 no incluye la carpeta de Hong Kong."),
    ("F22", "DC2015_JSON — elecciones de Consejo de Distrito 2015", "UnKnoWn-Consortium (GitHub)", "Plataforma colaborativa",
     "https://github.com/UnKnoWn-Consortium/DC2015_JSON", "JSON", "Descarga directa", "no declara", "Personas identificables", "Circunscripciones 2015",
     "2015", "Única vez", "Descartada", "Búsqueda web de GeoJSON de distritos", "Sin licencia y con datos de candidatos; límites de 2015 superados por F01."),
    ("F23", "hkdatasets — accidentes de tránsito 2014–2019", "Hong Kong Districts Info (paquete R)", "Academia u ONG",
     "https://github.com/Hong-Kong-Districts-Info/hkdatasets/raw/master/data-ready/hk_accidents.fst", "Otro", "Descarga directa", "MIT", "No",
     "Hong Kong (puntos)", "2014 a 2019", "Única vez", "Descartada", "La IA leyó el README del paquete",
     "Formato .fst de R no legible en el entorno; datos viejos. Útil para seguridad vial (wiki 4.9) en una siguiente entrega."),
    ("F24", "Temperatura media del último minuto por estación (CSV)", "Hong Kong Observatory (HKO)", "Entidad pública",
     "https://data.weather.gov.hk/weatherAPI/hko_data/regional-weather/latest_1min_temperature.csv", "CSV", "Bloquea scripts o pide sesión", DATAGOV,
     "No", "Estaciones HKO", "Tiempo real", "Tiempo real", "Candidata", "La IA la encontró en data.gov.hk",
     "robots.txt bloquea a la herramienta de la IA; probar desde un script propio para lecturas en vivo."),
    ("F25", "Population and Household Statistics Analysed by District Council District 2024 (PDF)", "Census and Statistics Department (C&SD)",
     "Entidad pública", "https://gia.info.gov.hk/general/202503/28/P2025032800255_490449_1_1743135159842.pdf", "PDF", "Solo PDF o página", DATAGOV,
     "Solo conteos por zona", "18 distritos", "2024", "Anual", "Descartada", "Búsqueda web",
     "La herramienta no pudo extraer las tablas del PDF; los mismos datos vienen en F03."),
    ("F26", "API de indicadores del Banco Mundial (Hong Kong)", "Banco Mundial", "Entidad pública",
     "https://api.worldbank.org/v2/country/HKG/indicator/NY.GDP.MKTP.CD?format=json", "JSON", "API sin llave", "no declara", "No", "Hong Kong",
     "1960 a 2025", "Anual", "No responde", "La IA la propuso para series macro",
     "El permiso para consultarla venció sin respuesta en la sesión; se reemplazó por FRED (F10–F12)."),
    ("F27", "Archivo histórico de DATA.GOV.HK (list-files)", "Digital Policy Office (DATA.GOV.HK)", "Entidad pública",
     "https://api.data.gov.hk/v1/historical-archive/list-files", "JSON", "API sin llave", DATAGOV, "No", "Hong Kong", "Varía", "Diaria",
     "Candidata", "La IA la probó para históricos de calidad del aire", "Responde, pero no se integró; sirve para series históricas de AQHI."),
    ("F28", "Página de referencia Cerebro Lima", "Proyecto de referencia (Vercel)", "No se sabe", "https://cerebro-lima.vercel.app/#panorama",
     "Página web", "Bloquea scripts o pide sesión", "no declara", "No", "Lima", "—", "No se sabe", "Descartada",
     "La entregó el usuario como modelo", "No es fuente de datos de HK; se usó solo como referencia de pestañas. El sitio bloquea acceso automatizado."),
]
FCOLS = ["id", "nombre", "entidad", "tipo", "url", "formato", "obtencion", "licencia", "personas", "cobertura", "periodo", "frecuencia",
         "estado", "como", "obs"]

# ---------------------------------------------------------------------------
# 3. Bitácora IA
# ---------------------------------------------------------------------------
BITACORA = [
    ("Leer la página de referencia Cerebro Lima para copiar su estructura de pestañas",
     "No pudo abrirla: el sitio bloquea acceso automatizado. Armó la estructura con los nombres de pestaña que le dimos",
     "Revisamos que las 10 pestañas del lago coinciden con las que pedimos", "Parcial", "F28",
     "Si la IA no puede ver la referencia, hay que darle la estructura explícita y confirmar su interpretación (p. ej. «atlas_distribuidos»)."),
    ("Conseguir los límites de los 18 distritos de Hong Kong", "Propuso geoBoundaries y la colección de mapas de Highcharts",
     "Probamos las rutas: geoBoundaries dio 404 y Highcharts no trae Hong Kong", "Inventado", "F20, F21",
     "La IA arma rutas «plausibles» que no existen; cada URL se prueba antes de anotarla."),
    ("Mismo encargo: límites de distritos", "Encontró el GeoJSON de circunscripciones DCCA 2019 de la EAC y disolvió los 18 distritos",
     "Contamos 452 circunscripciones y 18 distritos; comparamos áreas con Wikipedia", "Correcto", "F01",
     "Los polígonos administrativos incluyen mar: para densidad hay que usar el área terrestre oficial (F04)."),
    ("Traer el Censo 2021 por distrito", "Primero el XLSX/ZIP oficial (no lo pudo leer), luego la API get.php (error de parámetro) y al final el WFS de CSDI",
     "Sumamos grupos de edad y tipos de vivienda: igualan la población total en los 18 distritos", "Correcto", "F02, F18, F19",
     "Hubo que darle varias vueltas; el WFS fue la vía que funcionó. Pedir validaciones de suma es barato y atrapa errores."),
    ("Serie anual de ingreso y población por distrito", "Usó la Encuesta General de Hogares 2016–2025 vía CSDI",
     "La población 2021 de la serie coincide con el censo; ranking de ingreso igual al de la wiki", "Correcto", "F03",
     "Hay datos hasta 2025, más nuevos de lo que esperábamos."),
    ("Red de transporte público con coordenadas", "Encontró la compilación hk-bus-crawling con 15.287 paradas de todos los operadores",
     "Revisamos licencia (GPL-2.0) y que las paradas caigan en tierra: solo 3 se asignaron por cercanía", "Correcto", "F05",
     "Las compilaciones comunitarias ahorran mucho trabajo, pero su licencia debe anotarse."),
    ("Hospitales con urgencias", "Tomó nombres y coordenadas de un paquete npm (ane-hk)",
     "Revisamos que son los 18 hospitales A&E; no contrastamos coordenadas con una fuente oficial", "Parcial", "F06",
     "Falta el nuevo hospital de Kai Tak (abre en octubre de 2026, según prensa). Hay que buscar la lista oficial de la HA."),
    ("Estaciones meteorológicas con coordenadas", "Leyó la tabla del Observatorio (86 estaciones, en grados-minutos-segundos)",
     "Cruzamos con la capa CSDI: 24 de 26 a menos de 1 m; Tai Po a 300 m por reubicación", "Correcto", "F07, F08",
     "Tener dos fuentes de la misma cosa permite medir la calidad de las coordenadas."),
    ("Calidad del aire actual", "Leyó el RSS del AQHI de 18 estaciones",
     "La hora publicada (2026-10-06 07:30 HKT) parece de casi un día antes de la consulta", "Parcial", "F09",
     "Una captura puntual no sirve para tiempo real: hay que agendar la ingesta y guardar la hora."),
    ("Indicadores macro (PIB per cápita, inflación)", "Propuso la API del Banco Mundial; como no se pudo usar, tomó las series de FRED",
     "Revisamos que la serie constante coincide con la corriente en 2015 (base 2015, aunque FRED dice 2010)", "Parcial", "F10, F11, F12, F26",
     "Las etiquetas de las fuentes también pueden estar mal; un control simple lo detecta."),
    ("Delitos y seguridad", "Trajo el CSV mensual de la Policía y el informe Police in Figures 2025",
     "La suma mensual 2024 (94.948) no cuadra con el anual (94.747); ninguna fuente trae distritos", "Parcial", "F13, F14",
     "Anotar la diferencia en vez de forzar que cuadre; probablemente son revisiones."),
    ("Escucha: qué se dice de Hong Kong", "Hizo 8 búsquedas por tema, guardó 32 titulares y les calculó tono con un léxico",
     "Leímos los titulares: «Sham Shui Po needs more…» salía positivo por la palabra «more»; se corrigió el léxico", "Parcial", "F16, F17",
     "El tono por léxico es una señal gruesa; hay que revisarlo a mano y declararlo como heurístico."),
    ("Comparar el lago con la wiki del grupo", "Hizo tabla de cifras coincidentes y de cobertura por índice",
     "Recalculamos: delitos 2025, población de Nuevos Territorios y tenencia de vivienda coinciden; el % de 65+ difiere por definición", "Correcto", "F02, F13, F14",
     "Cifras distintas no siempre son errores: a veces cambian la base (con o sin empleadas domésticas)."),
    ("Llenar el diccionario de datos", "Generó el diccionario con un script que perfila cada columna (tipo, ejemplo, vacíos, rango)",
     "Revisamos una muestra de filas contra los CSV del lago", "Correcto", "",
     "Automatizar la parte medible del diccionario evita errores de copia; el significado y la pregunta sí requieren criterio."),
]

# ---------------------------------------------------------------------------
# 2. Diccionario
# ---------------------------------------------------------------------------
MAPA_F = {"eac_dcca_2019": "F01", "census2021_dc": "F02", "ghs_dc_2016_2025": "F03", "wiki_distritos": "F04",
          "hkbus_routefarelist": "F05", "ha_hospitales": "F06", "hko_estaciones": "F07", "hko_csdi_humedad": "F08",
          "epd_aqhi": "F09", "hkpf_delitos": "F13", "unemployment_recent": "F15", "noticias_escucha": "F16",
          "hkopendata_lugares": "F17"}
DERIVADO = {"atlas_zonas": "F01; F02; F03; F05; F06", "atlas_distribuidos": "F05; F06; F07", "correlaciones": "F02; F03; F05; F06"}
MACRO = {"anio": "F10; F11; F12", "pib_per_capita_usd_corrientes": "F10", "pib_per_capita_usd_constantes_2015": "F11",
         "inflacion_ipc_pct": "F12", "crecimiento_real_pib_pc_pct": "F11"}
FECHAS = {"publicado_hkt", "fecha_captura", "inicio_operacion", "periodo"}
COORD = {"lat", "lon", "centro_lat", "centro_lon"}


def id_fuente(fila):
    tabla, campo, fid = fila["Archivo o tabla en el lago"], fila["Nombre del campo en el lago"], fila["ID fuente"]
    if "macro_anual" in tabla:
        return MACRO[campo]
    if "indicadores_recientes" in tabla:
        return "F14; F15"
    if fid == "derivado":
        return DERIVADO[tabla.split("/")[1]]
    return "; ".join(dict.fromkeys(MAPA_F[x.strip()] for x in fid.split(";")))


def tipo(fila):
    campo, t, rol = fila["Nombre del campo en el lago"], fila["Tipo de dato"], fila["Clasificación"]
    if campo in COORD:
        return "Coordenada / geometría"
    if campo in FECHAS:
        return "Fecha"
    if rol.endswith("Identificador"):
        return "Código"
    return t


def diccionario():
    d = pd.read_csv(ROOT / "catalogo" / "diccionario_datos.csv", dtype=str).fillna("")
    d["ID fuente"] = d.apply(id_fuente, axis=1)
    d["Tipo de dato"] = d.apply(tipo, axis=1)
    d["Clasificación"] = "Pública"
    d.loc[d["Nombre del campo en el lago"] == "titular", "Clasificación"] = "Interna"
    d.loc[d["Nombre del campo en el lago"] == "titular", "Transformación que le hicieron"] += \
        " Clasificado «Interna» hasta revisar si nombra personas."

    # Claves de public/data/panorama.json que no existen en curated/
    p = json.loads((ROOT / "public" / "data" / "panorama.json").read_text())
    k, de = p["kpis"], p["destacados"]
    pan = [
        ("F02", "kpis.poblacion_censo_2021", "t_pop (suma)", "Población total de los 18 distritos en el Censo 2021.", "Entero", "personas", k["poblacion_censo_2021"], "≥ 0", "Suma de t_pop de los 18 distritos (excluye población marina).", "¿Cuánta gente vive en Hong Kong? (Wiki 1)"),
        ("F03", "kpis.poblacion_tierra_mitad_anio_2025", "my_lp (suma, 2025)", "Población terrestre de Hong Kong a mitad de 2025.", "Entero", "personas", k["poblacion_tierra_mitad_anio_2025"], "≥ 0", "Suma de my_lp 2025 de los 18 distritos.", "¿Cuánta gente vive hoy en Hong Kong? (Wiki 1)"),
        ("F02", "kpis.hogares_2021", "dh (suma)", "Total de hogares domésticos en 2021.", "Entero", "hogares", k["hogares_2021"], "≥ 0", "Suma de dh.", "¿Cuántos hogares hay en la ciudad?"),
        ("F04", "kpis.area_terrestre_km2", "Area (km²) (suma)", "Superficie terrestre total de los 18 distritos.", "Decimal", "km²", k["area_terrestre_km2"], "≥ 0", "Suma del área oficial por distrito.", "¿Qué tan grande es el territorio? (Wiki 6.5)"),
        ("F02; F04", "kpis.densidad_media_hab_km2", "(calculado)", "Densidad media de la ciudad.", "Entero", "hab/km²", int(k["densidad_media_hab_km2"]), "≥ 0", "Población 2021 total / área terrestre total.", "¿Qué tan densa es Hong Kong? (Wiki 1.2)"),
        ("F02", "kpis.pct_65_mas_2021", "age_5 / t_pop (sumas)", "Porcentaje de la población con 65 años o más en 2021.", "Decimal", "%", k["pct_65_mas_2021"], "0 a 100", "100 × suma age_5 / suma t_pop (incluye empleadas domésticas).", "¿Qué tan envejecida está la ciudad? (Wiki 1.3)"),
        ("F03", "kpis.ingreso_mediano_hogar_rango_distritos_2025", "ma_hh (mín. y máx., 2025)", "Ingreso mediano del hogar del distrito más pobre y del más rico en 2025.", "Entero", "HK$ por mes", "[24900, 45000]", "lista de 2 valores ≥ 0", "Mínimo y máximo de ma_hh 2025 entre distritos.", "¿Qué tan grande es la brecha de ingreso entre distritos? (Wiki 1.7)"),
        ("F10", "kpis.pib_per_capita_usd_2025", "PCAGDPHKA646NWDB (2025)", "PIB per cápita de 2025 en dólares corrientes.", "Decimal", "US$ corrientes", k["pib_per_capita_usd_2025"], "≥ 0", "Último año de la serie, redondeado.", "¿Qué tan rica es la economía? (Wiki 3.1)"),
        ("F12", "kpis.inflacion_ipc_2025_pct", "FPCPITOTLZGHKG (2025)", "Inflación al consumidor de 2025.", "Decimal", "%", k["inflacion_ipc_2025_pct"], "−5 a 10", "Último año de la serie, 2 decimales.", "¿Cómo está el costo de vida?"),
        ("F05", "kpis.paradas_transporte_publico", "(calculado)", "Número de paradas de transporte público con rutas.", "Entero", "paradas", k["paradas_transporte_publico"], "≥ 0", "Conteo de paradas con al menos una ruta.", "¿Qué tan extensa es la red de transporte?"),
        ("F05", "kpis.estaciones_mtr", "(calculado)", "Número de estaciones del MTR.", "Entero", "estaciones", k["estaciones_mtr"], "≥ 0", "Conteo de paradas del operador mtr.", "¿Qué tan extensa es la red de metro?"),
        ("F06", "kpis.hospitales_urgencias", "(calculado)", "Número de hospitales públicos con urgencias.", "Entero", "hospitales", k["hospitales_urgencias"], "≥ 0", "Conteo de filas de F06.", "¿Cuántas urgencias públicas hay? (Wiki 1.8)"),
        ("F07", "kpis.estaciones_meteorologicas", "(calculado)", "Número de estaciones del Observatorio.", "Entero", "estaciones", k["estaciones_meteorologicas"], "≥ 0", "Conteo de filas de F07.", "¿Qué tan monitoreado está el clima?"),
        ("F09", "kpis.estaciones_calidad_aire", "(calculado)", "Número de estaciones de calidad del aire.", "Entero", "estaciones", k["estaciones_calidad_aire"], "≥ 0", "Conteo de filas de F09.", "¿Qué tan monitoreado está el aire?"),
        ("F16", "kpis.noticias_monitoreadas", "(calculado)", "Número de titulares en la escucha.", "Entero", "titulares", k["noticias_monitoreadas"], "≥ 0", "Conteo de filas de F16.", "¿Cuánta conversación pública se está siguiendo?"),
    ]
    for clave, preg in [("mas_poblado", "¿Cuál es el distrito más poblado?"), ("mas_denso", "¿Cuál es el distrito más denso? (Wiki 1.2)"),
                        ("mayor_ingreso", "¿Cuál es el distrito más rico? (Wiki 1.7)"), ("menor_ingreso", "¿Cuál es el distrito más pobre? (Wiki 1.7)"),
                        ("mas_envejecido", "¿Cuál es el distrito más envejecido? (Wiki 1.3)"),
                        ("mayor_crecimiento_poblacion_2016_2025", "¿Qué distrito crece más? (Wiki 1.1)")]:
        pan.append(("F02; F03; F04", f"destacados.{clave}", "(calculado)", f"Nombre en español del distrito que encabeza «{clave.replace('_', ' ')}».",
                    "Texto", "—", de[clave], "Uno de los 18 nombres de distrito", "Distrito con el valor máximo (o mínimo) del indicador en el atlas.", preg))
    extra = pd.DataFrame([{
        "ID fuente": a, "Archivo o tabla en el lago": "public/data/panorama.json", "Nombre del campo en el lago": b,
        "Nombre original en la fuente": c, "Qué significa (en una frase)": e, "Tipo de dato": t, "Unidad de medida": u,
        "Ejemplo de valor": str(ej), "¿Puede venir vacío?": "No", "Valores permitidos o rango": r, "Clasificación": "Pública",
        "Transformación que le hicieron": tr, "Pregunta del sistema que ayuda a responder": q} for a, b, c, e, t, u, ej, r, tr, q in pan])
    return pd.concat([d, extra], ignore_index=True)


# ---------------------------------------------------------------------------
# Escritura en la plantilla
# ---------------------------------------------------------------------------
def preparar_filas(ws, n_filas, ultima_col):
    """Asegura formato amarillo y validaciones hasta la fila 4 + n_filas."""
    fin = 4 + n_filas
    if fin <= 44:
        return 44
    modelo = [ws.cell(row=5, column=c) for c in range(1, ultima_col + 1)]
    for r in range(45, fin + 1):
        for c, m in enumerate(modelo, 1):
            cel = ws.cell(row=r, column=c)
            cel.fill, cel.font, cel.border, cel.alignment, cel.number_format = copy(m.fill), copy(m.font), copy(m.border), copy(m.alignment), m.number_format
        if ws.row_dimensions[5].height:
            ws.row_dimensions[r].height = ws.row_dimensions[5].height
    for dv in ws.data_validations.dataValidation:
        nuevo = re.sub(r"([A-Z]+)5:([A-Z]+)44", lambda m: f"{m.group(1)}5:{m.group(2)}{fin}", str(dv.sqref))
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(nuevo)
    return fin


def escribir(ws, filas):
    for i, fila in enumerate(filas, 5):
        for j, v in enumerate(fila, 1):
            if v is None or (isinstance(v, float) and pd.isna(v)):
                v = ""
            ws.cell(row=i, column=j, value=v)


def main(plantilla, salida):
    wb = openpyxl.load_workbook(plantilla)
    ins = wb["Instrucciones"]
    ins["B5"] = "Ciudad elegida: Hong Kong (Región Administrativa Especial de Hong Kong, China)"
    ins["B6"] = "Integrantes: Manuel Arenas Lara"
    ins["B7"] = f"Enlace al repositorio o a la carpeta del lago: {REPO}"

    fu = wb["1. Fuentes"]
    fin_f = preparar_filas(fu, len(FUENTES), 16)
    escribir(fu, [list(f[:12]) + [FECHA] + list(f[12:]) for f in FUENTES])

    dic = diccionario()
    di = wb["2. Diccionario"]
    fin_d = preparar_filas(di, len(dic), 13)
    escribir(di, dic.values.tolist())

    bi = wb["3. Bitácora IA"]
    fin_b = preparar_filas(bi, len(BITACORA), 8)
    escribir(bi, [[i, "Claude (Anthropic)", *b] for i, b in enumerate(BITACORA, 1)])

    # El resumen debe contar todas las filas llenas
    hojas_fin = {"1. Fuentes": fin_f, "2. Diccionario": fin_d, "3. Bitácora IA": fin_b}
    res = wb["Resumen"]
    for row in res.iter_rows(min_col=2, max_col=2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                f = c.value
                for hoja, fin in hojas_fin.items():
                    f = re.sub(rf"('{re.escape(hoja)}'!\$?[A-Z]+\$?5:\$?[A-Z]+\$?)44", rf"\g<1>{fin}", f)
                c.value = f
    wb.save(salida)
    print(f"Fuentes: {len(FUENTES)} · Campos: {len(dic)} · Bitácora: {len(BITACORA)} · filas diccionario hasta {fin_d}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
