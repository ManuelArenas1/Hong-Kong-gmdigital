"""Exporta catalogo/diccionario_datos.csv a Excel con formato."""
import json
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
d = pd.read_csv(ROOT / "catalogo" / "diccionario_datos.csv", dtype=str).fillna("")
cat = json.loads((ROOT / "catalogo" / "catalogo.json").read_text())

wb = Workbook()
F = "Arial"
head_font = Font(name=F, bold=True, color="FFFFFF", size=10)
head_fill = PatternFill("solid", fgColor="1F3A5F")
body = Font(name=F, size=9)
thin = Side(style="thin", color="D0D5DD")
borde = Border(left=thin, right=thin, top=thin, bottom=thin)
zebra = PatternFill("solid", fgColor="F3F6FA")

def hoja(ws, df, anchos):
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append(list(r))
    for j, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for c in ws[1]:
        c.font, c.fill, c.border = head_font, head_fill, borde
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 34
    for i, row in enumerate(ws.iter_rows(min_row=2), 2):
        for c in row:
            c.font, c.border = body, borde
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if i % 2 == 0:
                c.fill = zebra
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

ws = wb.active
ws.title = "Diccionario"
hoja(ws, d, [22, 34, 30, 26, 46, 10, 18, 22, 10, 30, 18, 50, 46])

fu = pd.DataFrame([{
    "ID fuente": f["id"], "Dominio": f.get("dominio", ""), "Título": f.get("titulo", ""),
    "Proveedor": f.get("proveedor", ""), "URL": f.get("url") or "; ".join(f.get("urls", [])),
    "Licencia": f.get("licencia", ""), "Notas / validación": " ".join(x for x in [f.get("notas", ""), f.get("validacion", "")] if x),
} for f in cat["fuentes"]] + [{
    "ID fuente": "derivado", "Dominio": "atlas_zonas / atlas_distribuidos / correlaciones",
    "Título": "Campos calculados en el lago a partir de varias fuentes", "Proveedor": "Pipeline del lago (pipeline/build.py)",
    "URL": "", "Licencia": "Hereda la de sus fuentes", "Notas / validación": "Ver columna 'Transformación que le hicieron'."}])
hoja(wb.create_sheet("Fuentes"), fu, [22, 22, 46, 40, 50, 30, 60])
wb["Fuentes"].freeze_panes = "B2"

notas = wb.create_sheet("Notas")
txt = [
    ("Diccionario de datos — Data lake Cerebro HK", True),
    (f"{len(d)} campos de {d['Archivo o tabla en el lago'].nunique()} tablas de la capa curated/ (una fila por campo y tabla).", False),
    ("ID fuente: corresponde a la hoja 'Fuentes' y a catalogo/catalogo.json. 'derivado' = calculado en el lago; varias fuentes se separan con ';'.", False),
    ("Clasificación: nivel de sensibilidad + rol del campo. Todo el lago es 'Pública' (datos abiertos, sin datos personales). Rol: Identificador, Dimensión, Métrica, Geoespacial, Metadato o Texto.", False),
    ("Tipo de dato, Ejemplo de valor, ¿Puede venir vacío? y Valores permitidos o rango se calcularon automáticamente sobre los datos actuales.", False),
    ("Pregunta del sistema: entre paréntesis se indica el subíndice de la wiki 'Perfil de la ciudad de Hong Kong' al que aporta (p. ej. Wiki 1.7).", False),
    ("Se regenera con: python pipeline/diccionario.py && python pipeline/diccionario_xlsx.py", False),
]
for i, (t, b) in enumerate(txt, 1):
    c = notas.cell(row=i, column=1, value=t)
    c.font = Font(name=F, size=12 if b else 10, bold=b)
    c.alignment = Alignment(wrap_text=True, vertical="top")
notas.column_dimensions["A"].width = 120

wb.save(ROOT / "catalogo" / "diccionario_datos.xlsx")
print("ok")
